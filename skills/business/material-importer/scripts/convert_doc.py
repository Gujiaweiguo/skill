#!/usr/bin/env python3
"""溯源携带的 markitdown 包装器（pptx/docx/pdf → Markdown）。

背景：P0「原件为源契约」要求每个派生产物携带 provenance；裸
`uv run markitdown file > raw/x.md` 产出无溯源的裸 md。本脚本补上：
  1. 从输入文件向上找 config/company.yaml → 源记录为 base 相对 posix 路径；
     找不到则记绝对路径。
  2. sha256 源文件。
  3. 跑 markitdown（skill venv），失败打印 stderr 并 exit 1，不写部分产物。
  4. pptx 额外跑两项 advisory 检测（v2 起三态语义）：
     - pptx_table_guard.py detect 表格检测 → pptx_suspects 行
     - python-pptx 逐 slide text-sparse 扫描 → pptx_text_sparse 行
     各自三值：发现（含 slide 号与原因）/ none（仅=检测成功且无发现）/
     undetermined (detect-error: 原因)（检测失败，永不渲染成 none）。
     检测是 advisory：失败不影响文本转换产物，原因进信号行 + stderr。
  5. 输出 = provenance 注释块 + 转换内容；覆盖已有输出允许（重提取幂等）。

用法：uv run scripts/convert_doc.py <input.(pptx|docx|pdf)> <output.md>
（非 PEP 723 独立脚本 —— markitdown 在 skill venv pyproject 里。）
"""

from __future__ import annotations

import hashlib
import json
import subprocess
import sys
from datetime import date
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
except Exception:  # pragma: no cover — 缺 python-pptx 时 text-sparse 检测降级为 undetermined
    Presentation = None  # type: ignore[assignment]
    MSO_SHAPE_TYPE = None  # type: ignore[assignment]
    PP_PLACEHOLDER = None  # type: ignore[assignment]

EXTRACTOR_VERSION = 2
FIDELITY = "text=high; structure=variable (pptx 合并表格不可靠 — 见 pptx_table_guard)"
# text-sparse 页级判定常量（ppt-evidence-ruleset-v1；固定常量不建配置系统）
TEXT_SPARSE_MAX_VISIBLE_CHARS = 50  # 同 slide 正文可见字符（非空白）≤ 50
TEXT_SPARSE_MIN_PICS = 1            # 同 slide 图片形状 ≥ 1
RULESET_VERSION = "ppt-evidence-ruleset-v1"
SKILL_DIR = Path(__file__).resolve().parent.parent


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def resolve_source_relpath(path: Path) -> str:
    """向上找 config/company.yaml；命中 → base 相对 posix 路径，否则绝对路径。"""
    for parent in path.resolve().parents:
        if (parent / "config" / "company.yaml").is_file():
            try:
                return path.resolve().relative_to(parent).as_posix()
            except ValueError:
                break
    return str(path.resolve())


def run_markitdown(src: Path) -> str:
    proc = subprocess.run(
        ["uv", "run", "markitdown", str(src)],
        cwd=str(SKILL_DIR),
        capture_output=True,
        text=True,
    )
    if proc.returncode != 0:
        print(f"markitdown failed (rc={proc.returncode}):\n{proc.stderr}", file=sys.stderr)
        raise SystemExit(1)
    return proc.stdout


def _warn(what: str, reason: str) -> None:
    print(f"warning: {what} — 信号记为 undetermined（{reason}）", file=sys.stderr)


def detect_pptx_suspects(src: Path) -> tuple[str, list[dict], str | None]:
    """表格守卫 detect，三态：("ok", findings, None) / ("failed", [], 原因)。

    检测失败不伪装成「成功无发现」——none 与 undetermined 由渲染层区分。
    """
    try:
        proc = subprocess.run(
            ["uv", "run", "scripts/pptx_table_guard.py", "detect", str(src), "--json"],
            cwd=str(SKILL_DIR),
            capture_output=True,
            text=True,
        )
    except OSError as exc:
        _warn("pptx_table_guard 无法启动", str(exc))
        return "failed", [], f"spawn-error: {exc}"
    if proc.returncode != 0:
        lines = (proc.stderr or "").strip().splitlines()
        reason = lines[-1].strip() if lines else f"rc={proc.returncode}"
        _warn("pptx_table_guard detect 失败", reason)
        return "failed", [], f"guard-rc{proc.returncode}: {reason}"
    try:
        data = json.loads(proc.stdout)
    except json.JSONDecodeError as exc:
        _warn("pptx_table_guard 输出不可解析", str(exc))
        return "failed", [], f"json-unparseable: {exc}"
    suspects: list[dict] = []
    for res in data.get("suspects", []):
        for f in res.get("findings", []):
            suspects.append(f)
    return "ok", suspects, None


def _iter_shapes(shapes):
    """深度优先展开组合形状（group），其余形状原样产出。"""
    for shape in shapes:
        yield shape
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from _iter_shapes(shape.shapes)
        except Exception:
            pass  # group 展开失败不影响形状本身已被产出；提取异常由调用方整体判 failed


def _shape_visible_chars(shape) -> int:
    """单个形状的正文可见字符数（非空白字符）；slide number 占位符计 0。

    文本提取自身失败时异常上抛 → 由 detect_text_sparse 整体判 undetermined，
    绝不把提取失败的页判成 text-sparse。
    """
    try:
        if getattr(shape, "is_placeholder", False) and (
            shape.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
        ):
            return 0
    except Exception:
        pass
    if not getattr(shape, "has_text_frame", False):
        return 0
    text = shape.text_frame.text or ""
    return sum(1 for ch in text if not ch.isspace())


def _is_picture(shape) -> bool:
    try:
        if shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
            return True
    except Exception:
        pass
    try:
        if getattr(shape, "is_placeholder", False) and (
            shape.placeholder_format.type == PP_PLACEHOLDER.PICTURE
        ):
            return True
    except Exception:
        pass
    return False


def detect_text_sparse(src: Path) -> tuple[str, list[int], str | None]:
    """逐 slide text-sparse 扫描，三态：("ok", [slide 号...], None) / ("failed", [], 原因)。

    判定仅基于同一 slide：正文可见字符 ≤ TEXT_SPARSE_MAX_VISIBLE_CHARS 且
    图片形状 ≥ TEXT_SPARSE_MIN_PICS。不计 slide number 注释、图片文件名，
    与整份 PPT 文本总量无关。slide 边界或正文提取失败 → 整体 failed
    （undetermined），绝不把失败页判成 sparse。信号语义 = 建议回看原件该页；
    不识别图示/流程/业务含义，不触发 OCR，不做封面/目录页型豁免（v1 决策）。
    """
    if Presentation is None:
        return "failed", [], "python-pptx-unavailable"
    try:
        prs = Presentation(str(src))
        sparse: list[int] = []
        for idx, slide in enumerate(prs.slides, start=1):
            chars = 0
            pics = 0
            for shape in _iter_shapes(slide.shapes):
                chars += _shape_visible_chars(shape)
                if _is_picture(shape):
                    pics += 1
            if chars <= TEXT_SPARSE_MAX_VISIBLE_CHARS and pics >= TEXT_SPARSE_MIN_PICS:
                sparse.append(idx)
        return "ok", sparse, None
    except Exception as exc:
        _warn("pptx text-sparse 扫描失败", f"{type(exc).__name__}: {exc}")
        return "failed", [], f"slide-scan-error: {type(exc).__name__}: {exc}"


def render_suspects_line(outcome: tuple[str, list[dict], str | None]) -> str:
    state, suspects, error = outcome
    if state == "failed":
        return f"pptx_suspects: undetermined (detect-error: {error})"
    if not suspects:
        return "pptx_suspects: none"
    parts = [
        f"slide {f['slide']} ({', '.join(f['reasons'])})"
        for f in suspects
    ]
    return "pptx_suspects: " + ", ".join(parts)


def render_text_sparse_line(outcome: tuple[str, list[int], str | None]) -> str:
    state, slides, error = outcome
    if state == "failed":
        return f"pptx_text_sparse: undetermined (detect-error: {error})"
    if not slides:
        return "pptx_text_sparse: none"
    return "pptx_text_sparse: " + ", ".join(f"slide {n} (text-sparse)" for n in slides)


def _outcome_summary(outcome: tuple[str, list, str | None]) -> str:
    state, items, _ = outcome
    return "undetermined" if state == "failed" else str(len(items))


def main(argv: list[str]) -> int:
    if len(argv) != 2:
        print(f"用法: {argv[0]} <input.(pptx|docx|pdf)> <output.md>", file=sys.stderr)
        return 2
    src = Path(argv[0]).resolve()
    out = Path(argv[1])
    if not src.is_file():
        print(f"error: not found: {src}", file=sys.stderr)
        return 2
    suffix = src.suffix.lower()
    if suffix not in (".pptx", ".docx", ".pdf"):
        print(f"error: unsupported suffix {suffix!r} (expect .pptx/.docx/.pdf)", file=sys.stderr)
        return 2

    # markitdown 对垃圾字节可能按纯文本 rc=0 输出（实测），先做魔数校验拦截
    magic = src.open("rb").read(5)
    if suffix in (".pptx", ".docx") and not magic.startswith(b"PK"):
        print(f"error: {src} is not a valid OOXML (zip) file — refusing to convert", file=sys.stderr)
        return 1
    if suffix == ".pdf" and magic != b"%PDF-":
        print(f"error: {src} is not a valid PDF file — refusing to convert", file=sys.stderr)
        return 1

    source_rel = resolve_source_relpath(src)
    sha = file_sha256(src)

    content = run_markitdown(src)  # 失败即 exit 1，不写部分产物

    suspects_out: tuple[str, list[dict], str | None] = ("ok", [], None)
    sparse_out: tuple[str, list[int], str | None] = ("ok", [], None)
    if suffix == ".pptx":
        suspects_out = detect_pptx_suspects(src)
        sparse_out = detect_text_sparse(src)

    header_lines = [
        f"<!-- source: {source_rel} -->",
        f"<!-- source_sha256: {sha} -->",
        f"<!-- extracted_by: markitdown via convert_doc.py v{EXTRACTOR_VERSION} ({date.today().isoformat()}) -->",
        f"<!-- fidelity: {FIDELITY} -->",
    ]
    if suffix == ".pptx":
        header_lines.append(f"<!-- {render_suspects_line(suspects_out)} -->")
        header_lines.append(f"<!-- {render_text_sparse_line(sparse_out)} -->")
        header_lines.append(f"<!-- pptx_evidence_ruleset: {RULESET_VERSION} -->")

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("\n".join(header_lines) + "\n" + content, encoding="utf-8")

    summary = ""
    if suffix == ".pptx":
        summary = (
            f" suspects={_outcome_summary(suspects_out)}"
            f" text_sparse={_outcome_summary(sparse_out)}"
        )
    print(
        f"ok: source={source_rel} sha256={sha[:8]} "
        f"out={out} ({out.stat().st_size} bytes){summary}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
