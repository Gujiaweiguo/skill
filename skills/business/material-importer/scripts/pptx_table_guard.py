#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=0.6.21"]
# ///
# ─── How to run ───
# 1. Install uv: curl -LsSf https://astral.sh/uv/install.sh | sh
# 2. Detect: uv run pptx_table_guard.py detect <file-or-dir> [--json]
# 3. Escalate: uv run pptx_table_guard.py escalate <file.pptx> --slides 17,21 [--outdir DIR]
# ──────────────────
"""PPTX 复杂表格守卫：检测 markitdown 转换高风险表格 + 定向升级渲染。

背景（实测证据）：markitdown 转换 pptx 文本段落 0% 丢失，但复杂表格结构被
破坏——真实 21 列合并表头战略模板表格被转成 3 行错位文本，肉眼无法察觉。
架构决策：自动检测 + 定向升级（只渲染可疑幻灯片 + 本地 OCR），不做全量视觉。

detect 子命令（报告工具，exit 0 恒定）：
  对每个 .pptx（.ppt 为旧二进制格式，python-pptx 不支持 → 报告为 skipped），
  逐 slide 找 GraphicFrame 表格，命中任一规则即 SUSPECT：
    a. 存在合并单元格（<a:tc> 带属性 gridSpan>1 / rowSpan>1 / hMerge / vMerge，
       直接在 table._tbl XML 上做 namespace 搜索，不依赖 python-pptx 版本 API）
    b. 列数 > 10
    c. 首行（表头）空字符串占比 > 30%

escalate 子命令：
  只渲染指定 slide 为 PNG：soffice --headless --convert-to pdf（临时目录
  /tmp/opencode 下）→ pdftoppm -png -r 150 -f N -l N 逐页。pdftoppm 缺失时
  明确报错 exit 2，不半跑。默认 outdir 为 <file_dir>/_suspect/<stem>/
  （输出位置由调用方选择，本工具不做 incoming→raw 路径决策）。
  渲染后打印下一步指示：用 ocr_extract.py 做 OCR（GPU 运行时，本工具不代跑），
  并写 {outdir}/provenance.json（源路径 / sha256 / 渲染页 / 时间 / 工具版本）。
"""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

MERGE_ATTRS = ("gridSpan", "rowSpan", "hMerge", "vMerge")

COL_LIMIT = 10
HEADER_EMPTY_RATIO = 0.30


# ─── detect ───────────────────────────────────────────────────────────


def _find_tables(slide):
    """yield slide 中所有 GraphicFrame 表格 (shape, table)。"""
    for shape in slide.shapes:
        if shape.has_table:
            yield shape, shape.table


def _has_merged_cells(tbl) -> bool:
    """在 <a:tbl> XML 上搜索合并属性（版本无关，最稳）。"""
    for tc in tbl._tbl.iter(f"{{{A_NS}}}tc"):
        for attr in MERGE_ATTRS:
            val = tc.get(attr)
            if val is None:
                continue
            if attr in ("hMerge", "vMerge"):
                if val == "1" or val.lower() == "true":
                    return True
            else:  # gridSpan / rowSpan
                try:
                    if int(val) > 1:
                        return True
                except ValueError:
                    continue
    return False


def _header_empty_ratio(table) -> float:
    """首行空字符串单元格占比；空表返回 0。"""
    if len(table.rows) == 0:
        return 0.0
    header = [c.text for c in table.rows[0].cells]
    if not header:
        return 0.0
    empty = sum(1 for t in header if t.strip() == "")
    return empty / len(header)


def inspect_table(table) -> tuple[list[str], int, int]:
    """返回 (reasons, cols, rows)。reasons 非空 = SUSPECT。"""
    reasons: list[str] = []
    cols = len(table.columns)
    rows = len(table.rows)
    if _has_merged_cells(table):
        reasons.append("merged-cells")
    if cols > COL_LIMIT:
        reasons.append(f"cols>{COL_LIMIT}")
    if _header_empty_ratio(table) > HEADER_EMPTY_RATIO:
        reasons.append("header-empty")
    return reasons, cols, rows


def detect_file(path: Path) -> dict:
    """检查单个 pptx，返回 {file, findings[], error}。"""
    result: dict = {"file": str(path), "findings": [], "error": None}
    try:
        prs = Presentation(str(path))
    except Exception as exc:  # noqa: BLE001 — 报告工具不中断
        result["error"] = f"{type(exc).__name__}: {exc}"
        return result
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for tbl_idx, (_shape, table) in enumerate(_find_tables(slide)):
            reasons, cols, rows = inspect_table(table)
            if reasons:
                result["findings"].append(
                    {
                        "slide": slide_idx,
                        "table_index": tbl_idx,
                        "cols": cols,
                        "rows": rows,
                        "reasons": reasons,
                    }
                )
    return result


def iter_pptx(target: Path):
    """walk .pptx 文件，.ppt 报 skipped。yield (path|None, skipped_path|None)。"""
    if target.is_file():
        candidates = [target]
    else:
        candidates = sorted(
            p for p in target.rglob("*") if p.suffix.lower() in (".pptx", ".ppt")
        )
    for p in candidates:
        if p.suffix.lower() == ".ppt":
            yield None, p
        else:
            yield p, None


def cmd_detect(argv: list[str]) -> int:
    use_json = "--json" in argv
    args = [a for a in argv if a != "--json"]
    if not args:
        print("usage: pptx_table_guard.py detect <file-or-dir> [--json]", file=sys.stderr)
        return 2
    target = Path(args[0])
    if not target.exists():
        print(f"error: not found: {target}", file=sys.stderr)
        return 2

    findings: list[dict] = []
    skipped: list[str] = []
    for pptx, ppt in iter_pptx(target):
        if ppt is not None:
            skipped.append(str(ppt))
            continue
        res = detect_file(pptx)
        if res["findings"] or res["error"]:
            findings.append(res)

    if use_json:
        print(json.dumps({"suspects": findings, "skipped": skipped}, ensure_ascii=False, indent=2))
    else:
        for res in findings:
            if res["error"]:
                print(f"{res['file']} :: ERROR :: {res['error']}")
            for f in res["findings"]:
                print(
                    f"{res['file']} :: slide {f['slide']} :: table idx {f['table_index']}"
                    f" :: {', '.join(f['reasons'])} ({f['cols']}x{f['rows']})"
                )
        if skipped:
            for s in skipped:
                print(f"{s} :: skipped :: legacy .ppt unsupported")
        total = sum(len(r["findings"]) for r in findings)
        n_files = len(findings)
        print(f"summary: {total} suspect table(s) in {n_files} file(s); {len(skipped)} skipped .ppt")
    return 0


# ─── escalate ─────────────────────────────────────────────────────────


def _tool_version(cmd: list[str]) -> str:
    try:
        out = subprocess.run(cmd, capture_output=True, text=True, timeout=30)
        return (out.stdout or out.stderr).strip().splitlines()[0][:120]
    except Exception:  # noqa: BLE001
        return "unknown"


def cmd_escalate(argv: list[str]) -> int:
    file_arg: str | None = None
    slides_arg: str | None = None
    outdir_arg: str | None = None
    i = 0
    while i < len(argv):
        a = argv[i]
        if a == "--slides":
            slides_arg = argv[i + 1] if i + 1 < len(argv) else None
            i += 2
        elif a == "--outdir":
            outdir_arg = argv[i + 1] if i + 1 < len(argv) else None
            i += 2
        elif a.startswith("--"):
            print(f"error: unknown option {a}", file=sys.stderr)
            return 2
        else:
            file_arg = a
            i += 1
    if not file_arg or not slides_arg:
        print(
            "usage: pptx_table_guard.py escalate <file.pptx> --slides 17,21 [--outdir DIR]",
            file=sys.stderr,
        )
        return 2
    src = Path(file_arg).resolve()
    if not src.is_file() or src.suffix.lower() != ".pptx":
        print(f"error: not a .pptx file: {src}", file=sys.stderr)
        return 2
    try:
        slides = sorted({int(s) for s in slides_arg.split(",") if s.strip()})
    except ValueError:
        print(f"error: bad --slides value: {slides_arg!r} (expect e.g. 17,21)", file=sys.stderr)
        return 2
    if not slides:
        print("error: --slides is empty", file=sys.stderr)
        return 2

    if shutil.which("soffice") is None:
        print("error: missing binary 'soffice' (LibreOffice) — required for escalate", file=sys.stderr)
        return 2
    if shutil.which("pdftoppm") is None:
        print("error: missing binary 'pdftoppm' (poppler-utils) — required for escalate", file=sys.stderr)
        return 2

    outdir = Path(outdir_arg).resolve() if outdir_arg else src.parent / "_suspect" / src.stem
    outdir.mkdir(parents=True, exist_ok=True)

    sha256 = hashlib.sha256(src.read_bytes()).hexdigest()
    with tempfile.TemporaryDirectory(prefix="pptx_guard_", dir="/tmp/opencode") as tmp:
        tmpdir = Path(tmp)
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmpdir), str(src)],
            check=True, capture_output=True, timeout=300,
        )
        pdf = tmpdir / (src.stem + ".pdf")
        if not pdf.exists():
            print(f"error: soffice produced no pdf for {src}", file=sys.stderr)
            return 2
        rendered = []
        for n in slides:
            subprocess.run(
                ["pdftoppm", "-png", "-r", "150", "-f", str(n), "-l", str(n),
                 str(pdf), str(outdir / f"slide_{n}")],
                check=True, capture_output=True, timeout=120,
            )
            png = outdir / f"slide_{n}-{n:02d}.png"
            if not png.exists():
                # pdftoppm zero-pads by total page count; find the single match
                cands = list(outdir.glob(f"slide_{n}-*.png"))
                if not cands:
                    print(f"error: no PNG rendered for slide {n}", file=sys.stderr)
                    return 2
                png = cands[0]
            rendered.append(n)

    provenance = {
        "source_path": str(src),
        "source_sha256": sha256,
        "rendered_slides": rendered,
        "rendered_at": datetime.now(timezone.utc).isoformat(),
        "tool_versions": {
            "soffice": _tool_version(["soffice", "--version"]),
            "pdftoppm": _tool_version(["pdftoppm", "-v"]),
        },
    }
    (outdir / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(f"rendered slides {rendered} -> {outdir}")
    print(f"next step: uv run scripts/ocr_extract.py {outdir} --output-dir {outdir}")
    return 0


# ─── main ─────────────────────────────────────────────────────────────


def main(argv: list[str]) -> int:
    if len(argv) < 2 or argv[1] not in ("detect", "escalate"):
        print(__doc__.splitlines()[0] if __doc__ else "", file=sys.stderr)
        print(
            "usage: pptx_table_guard.py <detect|escalate> ...\n"
            "  detect <file-or-dir> [--json]\n"
            "  escalate <file.pptx> --slides 17,21 [--outdir DIR]",
            file=sys.stderr,
        )
        return 2
    if argv[1] == "detect":
        return cmd_detect(argv[2:])
    return cmd_escalate(argv[2:])


if __name__ == "__main__":
    sys.exit(main(sys.argv))
