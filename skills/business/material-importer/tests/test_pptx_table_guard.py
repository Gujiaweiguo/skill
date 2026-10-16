"""pptx_table_guard 检测逻辑回归测试（不依赖 soffice/pdftoppm）。

detect 三条 SUSPECT 规则（合并单元格 / 列数>10 / 表头空占比>30%）+
.ppt 跳过 + JSON 输出格式。escalate 只测参数校验与缺二进制路径。
"""

import json
import os
import sys
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

SCRIPTS_DIR = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS_DIR))

import pptx_table_guard as guard  # noqa: E402


def make_pptx_with_table(path: Path, header: list[str], n_data_rows: int = 2) -> Path:
    """构建一个含单表格的 pptx：首行 header，后续 n_data 行全填充。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    n_cols = len(header)
    gf = slide.shapes.add_table(
        n_data_rows + 1, n_cols, Inches(0.3), Inches(0.3), Inches(9), Inches(3)
    )
    table = gf.table
    for ci, text in enumerate(header):
        table.cell(0, ci).text = text
    for ri in range(1, n_data_rows + 1):
        for ci in range(n_cols):
            table.cell(ri, ci).text = f"v{ri}-{ci}"
    prs.save(str(path))
    return path


class DetectRulesTest(unittest.TestCase):
    def setUp(self) -> None:
        import tempfile

        self.tmp = Path(tempfile.mkdtemp(prefix="ptg_test_", dir="/tmp/opencode"))

    def tearDown(self) -> None:
        import shutil

        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_normal_table_not_suspect(self) -> None:
        p = make_pptx_with_table(self.tmp / "normal.pptx", ["A", "B", "C"])
        res = guard.detect_file(p)
        self.assertEqual(res["findings"], [])
        self.assertIsNone(res["error"])

    def test_merged_cells_suspect(self) -> None:
        p = make_pptx_with_table(self.tmp / "merged.pptx", ["A", "B", "C"])
        prs = Presentation(str(p))
        table = next(s for s in prs.slides[0].shapes if s.has_table).table
        merged = False
        if hasattr(table.cell(0, 0), "merge"):
            table.cell(0, 0).merge(table.cell(0, 1))
            merged = True
        else:
            # 回退：直接在 XML 上设 gridSpan（测试自身也走 XML 层，与检测器同源规则）
            from pptx.oxml.ns import qn

            tcs = table._tbl.findall(qn("a:tr"))[0].findall(qn("a:tc"))
            tcs[0].set("gridSpan", "2")
            merged = True
        prs.save(str(p))
        self.assertTrue(merged)
        res = guard.detect_file(p)
        reasons = [r for f in res["findings"] for r in f["reasons"]]
        self.assertIn("merged-cells", reasons)

    def test_wide_table_suspect(self) -> None:
        p = make_pptx_with_table(
            self.tmp / "wide.pptx", [f"H{i}" for i in range(12)]
        )
        res = guard.detect_file(p)
        self.assertEqual(len(res["findings"]), 1)
        self.assertIn("cols>10", res["findings"][0]["reasons"])
        self.assertEqual(res["findings"][0]["cols"], 12)

    def test_header_empty_suspect(self) -> None:
        # 4/5 空 = 80% > 30%
        p = make_pptx_with_table(self.tmp / "hdr.pptx", ["", "", "", "", "X"])
        res = guard.detect_file(p)
        reasons = [r for f in res["findings"] for r in f["reasons"]]
        self.assertIn("header-empty", reasons)

    def test_ppt_skipped_in_dir_mode(self) -> None:
        make_pptx_with_table(self.tmp / "ok.pptx", ["A", "B"])
        (self.tmp / "legacy.ppt").write_bytes(b"\xd0\xcf\x11\xe0fake")
        collected_pptx, skipped = [], []
        for pptx, ppt in guard.iter_pptx(self.tmp):
            if ppt is not None:
                skipped.append(ppt.name)
            else:
                collected_pptx.append(pptx.name)
        self.assertEqual(collected_pptx, ["ok.pptx"])
        self.assertEqual(skipped, ["legacy.ppt"])

    def test_json_mode_well_formed(self) -> None:
        import contextlib
        import io

        # cmd_detect 目标目录与 setUp 的 tmp 不同即可（--json 校验输出结构）
        import tempfile

        run_tmp = Path(tempfile.mkdtemp(prefix="ptg_json_", dir="/tmp/opencode"))
        make_pptx_with_table(run_tmp / "wide.pptx", [f"H{i}" for i in range(12)])
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = guard.cmd_detect([str(run_tmp), "--json"])
        self.assertEqual(rc, 0)
        data = json.loads(buf.getvalue())
        self.assertIn("suspects", data)
        self.assertIn("skipped", data)
        flat = [f for r in data["suspects"] for f in r["findings"]]
        self.assertTrue(any("cols>10" in f["reasons"] for f in flat))


class EscalateArgTest(unittest.TestCase):
    """escalate 只测参数校验；不渲染（需 soffice/pdftoppm）。"""

    def test_missing_args_usage(self) -> None:
        self.assertEqual(guard.cmd_escalate([]), 2)

    def test_bad_slides_value(self) -> None:
        import tempfile

        with tempfile.NamedTemporaryFile(suffix=".pptx", dir="/tmp/opencode", delete=False) as f:
            f.write(b"fake")
            p = f.name
        try:
            self.assertEqual(guard.cmd_escalate([p, "--slides", "a,b"]), 2)
        finally:
            os.unlink(p)

    def test_missing_binary_path(self) -> None:
        # 空技巧：用不存在的文件先触发文件校验不合适；直接造合法 pptx，
        # stub PATH 使 soffice/pdftoppm 不可见 → 期望 exit 2 且报缺二进制。
        import tempfile

        tmp = Path(tempfile.mkdtemp(prefix="ptg_esc_", dir="/tmp/opencode"))
        p = make_pptx_with_table(tmp / "x.pptx", ["A", "B"])
        old_path = os.environ.get("PATH")
        os.environ["PATH"] = str(tmp)  # 空目录：找不到 soffice / pdftoppm
        try:
            import contextlib
            import io

            buf = io.StringIO()
            with contextlib.redirect_stderr(buf):
                rc = guard.cmd_escalate([str(p), "--slides", "1"])
            self.assertEqual(rc, 2)
            self.assertIn("missing binary", buf.getvalue())
        finally:
            os.environ["PATH"] = old_path or ""


if __name__ == "__main__":
    unittest.main()
