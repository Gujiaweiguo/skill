"""tests for scripts/convert_doc.py — 溯源携带的 markitdown 包装器。

两类用例：
1. subprocess 端到端（与真实调用方式一致）——provenance 头、pptx 信号行、失败退出码。
2. in-process（importlib 加载脚本 + 包装其 subprocess.run / Presentation）——
   检测三态语义的失败注入；跨子进程边界无法 monkeypatch，改在进程内替换依赖。

text-sparse 误报记录（ppt-evidence-ruleset-v1）：封面/目录/章节分隔页（少文有图）
按规则会被标记，属预期误报；用例断言信号仅为复核提示（rc=0、产物完整、无 OCR/识别
声明），实现不做页型豁免（v1 决策，见 SKILL.md 设计决策）。
"""

from __future__ import annotations

import contextlib
import hashlib
import importlib.util
import io
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from pptx import Presentation
from pptx.util import Inches

SKILL_DIR = Path(__file__).resolve().parent.parent
SCRIPT = SKILL_DIR / "scripts" / "convert_doc.py"


def run_script(src: Path, out: Path):
    return subprocess.run(
        [sys.executable, str(SCRIPT), str(src), str(out)],
        cwd=str(SKILL_DIR),
        capture_output=True,
        text=True,
        timeout=300,
    )


_module_seq = 0


def load_convert_doc():
    """进程内加载 convert_doc.py（独立模块名避免装载冲突）。"""
    global _module_seq
    _module_seq += 1
    spec = importlib.util.spec_from_file_location(f"convert_doc_ut_{_module_seq}", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _tiny_png(workdir: Path, name: str = "dot.png") -> Path:
    from PIL import Image

    p = workdir / name
    Image.new("RGB", (2, 2), color=(180, 180, 180)).save(p, format="PNG")
    return p


def _make_sparse_vs_dense_pptx(path: Path):
    """slide1 = 图 + 2 字（命中）；slide2 = 图 + 80 字正文（自身充足，不命中）。"""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    img = _tiny_png(path.parent)
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(str(img), Inches(1), Inches(1))
    tb1 = s1.shapes.add_textbox(Inches(0.5), Inches(3), Inches(4), Inches(0.5))
    tb1.text_frame.text = "目录"
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_picture(str(img), Inches(1), Inches(1))
    tb2 = s2.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(8), Inches(2))
    tb2.text_frame.text = "正文内容" * 20  # 80 个非空白字符 > 50
    prs.save(str(path))


def _make_simple_pptx(path: Path):
    """最小 pptx：1 图 + 2 字 → text-sparse 命中；无表格 → suspects 无命中。"""
    prs = Presentation()
    blank = prs.slide_layouts[6]
    img = _tiny_png(path.parent)
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(img), Inches(1), Inches(1))
    tb = s.shapes.add_textbox(Inches(0.5), Inches(3), Inches(4), Inches(0.5))
    tb.text_frame.text = "目录"
    prs.save(str(path))


class ConvertDocDocxTest(unittest.TestCase):
    def test_docx_provenance_header_and_body(self):
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            (root / "config").mkdir()
            (root / "config" / "company.yaml").write_text("company: test\n", encoding="utf-8")
            docx_path = root / "测试文档.docx"
            doc = Document()
            doc.add_paragraph("溯源契约测试正文内容")
            doc.save(str(docx_path))

            out = root / "out" / "测试文档.docx.md"
            proc = run_script(docx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            lines = text.splitlines()
            self.assertTrue(lines[0].startswith("<!-- source: "), lines[0])
            # base 相对路径（config/company.yaml 存在于同根）
            self.assertIn("source: 测试文档.docx -->", lines[0])
            self.assertEqual(
                lines[1],
                f"<!-- source_sha256: {hashlib.sha256(docx_path.read_bytes()).hexdigest()} -->",
            )
            self.assertTrue(lines[2].startswith("<!-- extracted_by: markitdown via convert_doc.py v"), lines[2])
            self.assertTrue(lines[3].startswith("<!-- fidelity: "), lines[3])
            # docx 无 pptx 信号行（suspects 与 text-sparse 均不出现）
            self.assertNotIn("pptx_suspects", text)
            self.assertNotIn("pptx_text_sparse", text)
            self.assertIn("溯源契约测试正文内容", text)


class ConvertDocPptxTest(unittest.TestCase):
    def _make_pptx(self, path: Path, cols: int):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        shape = slide.shapes.add_table(2, cols, Inches(0.5), Inches(0.5), Inches(9), Inches(1))
        table = shape.table
        for r in range(2):
            for c in range(cols):
                table.cell(r, c).text = f"R{r}C{c}"
        prs.save(str(path))

    def test_pptx_wide_table_flagged(self):
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            pptx_path = root / "wide.pptx"
            self._make_pptx(pptx_path, cols=11)

            out = root / "wide.pptx.md"
            proc = run_script(pptx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            self.assertIn("pptx_suspects:", text)
            self.assertIn("slide 1", text)
            self.assertIn("cols>10", text)

    def test_pptx_clean_table_none(self):
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            pptx_path = root / "clean.pptx"
            self._make_pptx(pptx_path, cols=4)

            out = root / "clean.pptx.md"
            proc = run_script(pptx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            # 回归锁定：检测成功且无发现仍是字面 none（下游兼容）
            self.assertIn("pptx_suspects: none", text)
            # 无图片 → text-sparse 也不命中（图片 ≥1 才算）
            self.assertIn("pptx_text_sparse: none", text)


class ConvertDocPptxTextSparseTest(unittest.TestCase):
    """text-sparse 页级判定（ppt-evidence-ruleset-v1）：同 slide 文本 ≤50 字且 ≥1 图。"""

    def test_sparse_page_flagged_and_dense_page_unaffected(self):
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            pptx_path = root / "mix.pptx"
            _make_sparse_vs_dense_pptx(pptx_path)

            out = root / "mix.pptx.md"
            proc = run_script(pptx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            # 页级判定：slide 1 命中（同 slide 图 + ≤50 字），信号含页码
            self.assertIn("pptx_text_sparse: slide 1 (text-sparse)", text)
            # slide 2 自身正文充足：不因整份文档文本量或其他页稀疏被标
            self.assertNotIn("slide 2 (text-sparse)", text)
            # suspects 与 text-sparse 相互独立
            self.assertIn("pptx_suspects: none", text)
            self.assertIn("pptx_evidence_ruleset: ppt-evidence-ruleset-v1", text)

    def test_text_only_slide_without_picture_not_flagged(self):
        """少文本但无图片 → 不命中（两条件必须同时满足）。"""
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            prs = Presentation()
            s = prs.slides.add_slide(prs.slide_layouts[6])
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(4), Inches(0.5))
            tb.text_frame.text = "纯文字短页"
            pptx_path = root / "textonly.pptx"
            prs.save(str(pptx_path))

            out = root / "textonly.pptx.md"
            proc = run_script(pptx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            self.assertIn("pptx_text_sparse: none", text)

    def test_cover_style_page_flagged_is_review_hint_only(self):
        """误报记录：封面/目录型（少文有图）被标记属预期误报，v1 不做页型豁免。

        断言信号仅为复核提示：转换 rc=0 不阻断、正文完整保留、头部无 OCR/识别声明。
        """
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            prs = Presentation()
            blank = prs.slide_layouts[6]
            img = _tiny_png(root, "cover.png")
            s = prs.slides.add_slide(blank)
            s.shapes.add_picture(str(img), Inches(2), Inches(2), width=Inches(4))
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(1))
            tb.text_frame.text = "产品介绍"
            pptx_path = root / "cover.pptx"
            prs.save(str(pptx_path))

            out = root / "cover.pptx.md"
            proc = run_script(pptx_path, out)
            self.assertEqual(proc.returncode, 0, proc.stderr)

            text = out.read_text(encoding="utf-8")
            self.assertIn("pptx_text_sparse: slide 1 (text-sparse)", text)
            self.assertIn("产品介绍", text)  # 转换产物照常写出，不被信号影响
            header = "\n".join(text.splitlines()[:7])
            self.assertNotIn("ocr", header.lower())  # 不自动 OCR
            self.assertNotIn("识别", header)  # 不声称已识别图示/流程/业务含义


class ConvertDocDetectionStateTest(unittest.TestCase):
    """检测三态语义：成功有发现 / 成功无发现(none) / 失败(undetermined，永不 none)。"""

    def _run_main(self, root: Path, guard=None, patch_presentation=None):
        mod = load_convert_doc()
        real_run = subprocess.run

        def wrapped_run(cmd, **kwargs):
            if guard is not None and any("pptx_table_guard" in str(c) for c in cmd):
                return guard(cmd)
            return real_run(cmd, **kwargs)

        mod.subprocess.run = wrapped_run
        if patch_presentation is not None:
            mod.Presentation = patch_presentation
        pptx_path = root / "in.pptx"
        _make_simple_pptx(pptx_path)
        out = root / "in.pptx.md"
        stdout_buf, stderr_buf = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(stdout_buf), contextlib.redirect_stderr(stderr_buf):
            # main(argv) 契约 = sys.argv[1:]，即 [src, out]
            rc = mod.main([str(pptx_path), str(out)])
        text = out.read_text(encoding="utf-8") if out.exists() else ""
        return rc, text, stdout_buf.getvalue(), stderr_buf.getvalue()

    def test_guard_failure_is_undetermined_not_none(self):
        def guard(cmd):
            return subprocess.CompletedProcess(cmd, returncode=1, stdout="", stderr="guard exploded")

        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            rc, text, stdout, stderr = self._run_main(Path(tmp), guard=guard)
            # advisory：检测失败不阻断转换产物
            self.assertEqual(rc, 0)
            self.assertTrue(text, "转换产物必须照常写出")
            self.assertIn("pptx_suspects: undetermined (detect-error: guard-rc1", text)
            self.assertNotIn("pptx_suspects: none", text)  # 失败不得伪装成无发现
            # 两信号相互独立：suspects 失败不影响 text-sparse 真实判定
            self.assertIn("pptx_text_sparse: slide 1 (text-sparse)", text)
            self.assertIn("undetermined", stderr)  # 失败原因同步进 stderr

    def test_guard_json_unparseable_is_undetermined(self):
        def guard(cmd):
            return subprocess.CompletedProcess(cmd, returncode=0, stdout="<not json>", stderr="")

        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            rc, text, stdout, stderr = self._run_main(Path(tmp), guard=guard)
            self.assertEqual(rc, 0)
            self.assertIn("pptx_suspects: undetermined (detect-error: json-unparseable", text)
            self.assertNotIn("pptx_suspects: none", text)

    def test_guard_success_empty_findings_still_renders_none(self):
        def guard(cmd):
            return subprocess.CompletedProcess(
                cmd, returncode=0, stdout=json.dumps({"suspects": []}), stderr=""
            )

        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            rc, text, stdout, stderr = self._run_main(Path(tmp), guard=guard)
            self.assertEqual(rc, 0)
            self.assertIn("pptx_suspects: none", text)  # none 仅=成功且无发现
            self.assertIn("pptx_text_sparse: slide 1 (text-sparse)", text)
            # stdout 汇总区分计数与 undetermined
            self.assertIn("suspects=0", stdout)
            self.assertIn("text_sparse=1", stdout)

    def test_slide_scan_failure_is_undetermined_never_sparse(self):
        def boom(*args, **kwargs):
            raise RuntimeError("pptx broken mid-scan")

        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            rc, text, stdout, stderr = self._run_main(Path(tmp), patch_presentation=boom)
            # 正文提取/边界枚举失败 → undetermined，绝不判 sparse；转换照常成功
            self.assertEqual(rc, 0)
            self.assertIn(
                "pptx_text_sparse: undetermined (detect-error: slide-scan-error: RuntimeError", text
            )
            self.assertNotIn("slide 1 (text-sparse)", text)
            # suspects 不受影响（有效 pptx 上真实 guard 无命中）
            self.assertIn("pptx_suspects: none", text)

    def test_stdout_summary_marks_undetermined(self):
        def guard(cmd):
            return subprocess.CompletedProcess(cmd, returncode=1, stdout="", stderr="x")

        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            rc, text, stdout, stderr = self._run_main(Path(tmp), guard=guard)
            self.assertIn("suspects=undetermined", stdout)
            self.assertIn("text_sparse=1", stdout)


class TextSparseHelperUnitTest(unittest.TestCase):
    """字符计数/图片判定 helper：slide number 注释、空白、非文本形状不计。"""

    def test_slide_number_placeholder_and_whitespace_not_counted(self):
        mod = load_convert_doc()
        slide_number_ph = SimpleNamespace(
            is_placeholder=True,
            placeholder_format=SimpleNamespace(type=mod.PP_PLACEHOLDER.SLIDE_NUMBER),
            has_text_frame=True,
            text_frame=SimpleNamespace(text="‹#›"),
        )
        self.assertEqual(mod._shape_visible_chars(slide_number_ph), 0)
        normal = SimpleNamespace(
            is_placeholder=False,
            has_text_frame=True,
            text_frame=SimpleNamespace(text="a b\nc"),
        )
        self.assertEqual(mod._shape_visible_chars(normal), 3)  # 只计非空白可见字符
        non_text = SimpleNamespace(is_placeholder=False, has_text_frame=False)
        self.assertEqual(mod._shape_visible_chars(non_text), 0)

    def test_picture_detection(self):
        mod = load_convert_doc()
        pic = SimpleNamespace(shape_type=mod.MSO_SHAPE_TYPE.PICTURE)
        self.assertTrue(mod._is_picture(pic))
        textbox = SimpleNamespace(shape_type=mod.MSO_SHAPE_TYPE.TEXT_BOX, is_placeholder=False)
        self.assertFalse(mod._is_picture(textbox))


class ConvertDocFailureTest(unittest.TestCase):
    def test_corrupt_input_exit1_no_output(self):
        with tempfile.TemporaryDirectory(dir="/tmp/opencode") as tmp:
            root = Path(tmp)
            bad = root / "corrupt.pptx"
            bad.write_bytes(b"not a real pptx file at all")
            out = root / "corrupt.pptx.md"

            proc = run_script(bad, out)
            self.assertEqual(proc.returncode, 1)
            self.assertFalse(out.exists(), "失败时不得写出部分产物")


if __name__ == "__main__":
    unittest.main()
